"""Leitura e escrita das planilhas (.xlsx / .csv), preservando o original.

A planilha de entrada nunca e alterada: o motor escreve sempre em uma copia,
mantendo formatacao, abas e colunas que nao entende - o arquivo continua
valido para reenvio ao Mercado Livre.
"""

from __future__ import annotations

import csv
import io
from abc import ABC, abstractmethod
from copy import copy
from decimal import Decimal
from pathlib import Path

EXTENSOES_EXCEL = {".xlsx", ".xlsm"}
EXTENSOES_TEXTO = {".csv", ".txt", ".tsv"}


class FormatoNaoSuportado(ValueError):
    pass


def _validar_extensao(caminho: Path) -> str:
    sufixo = caminho.suffix.lower()
    if sufixo in EXTENSOES_EXCEL or sufixo in EXTENSOES_TEXTO:
        return sufixo
    if sufixo == ".xls":
        raise FormatoNaoSuportado(
            f"{caminho.name}: formato .xls antigo nao e suportado. "
            "Abra no Excel/LibreOffice e salve como .xlsx."
        )
    raise FormatoNaoSuportado(
        f"{caminho.name}: extensao '{sufixo}' nao suportada (use .xlsx, .xlsm, .csv ou .tsv)"
    )


def _decodificar(bruto: bytes) -> str:
    """Exports do ML aparecem em UTF-8 (com ou sem BOM) e em Latin-1."""
    for codec in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return bruto.decode(codec)
        except UnicodeDecodeError:
            continue
    return bruto.decode("latin-1", errors="replace")


def _dialeto(amostra: str) -> csv.Dialect | type[csv.Dialect]:
    try:
        return csv.Sniffer().sniff(amostra, delimiters=";,\t|")
    except csv.Error:
        # Sem pistas: ponto-e-virgula e o padrao de CSV brasileiro.
        return csv.excel_tab if "\t" in amostra else _PontoEVirgula


class _PontoEVirgula(csv.Dialect):
    delimiter = ";"
    quotechar = '"'
    doublequote = True
    skipinitialspace = False
    lineterminator = "\r\n"
    quoting = csv.QUOTE_MINIMAL


class Tabela(ABC):
    """Planilha aberta para leitura das linhas e escrita do resultado."""

    linhas: list[list[object]]

    @abstractmethod
    def escrever(self, linha: int, coluna: int, valor: object) -> None:
        """Grava em coordenadas 0-based da matriz lida."""

    @abstractmethod
    def acrescentar_coluna(self, titulo: str, linha_cabecalho: int) -> int:
        """Cria uma coluna no fim da tabela e devolve o indice dela."""

    @abstractmethod
    def salvar(self, destino: str | Path) -> Path:
        """Escreve a copia preenchida."""

    def valor(self, linha: int, coluna: int | None) -> object:
        if coluna is None or linha >= len(self.linhas):
            return None
        celulas = self.linhas[linha]
        return celulas[coluna] if coluna < len(celulas) else None

    def cor_de_fundo(self, linha: int, coluna: int) -> str | None:
        """Cor de preenchimento da celula - o ML usa cor para marcar o que e editavel."""
        return None

    def opcoes_validacao(self, linha: int, coluna: int) -> list[str] | None:
        """Opcoes da lista suspensa da celula, quando existir."""
        return None


class TabelaExcel(Tabela):
    def __init__(self, caminho: Path, aba: str | None = None):
        from openpyxl import load_workbook

        self.caminho = caminho
        # Duas leituras: uma com os valores ja calculados (para decidir) e outra
        # preservando formulas e estilos (para salvar sem estragar o arquivo).
        self._wb = load_workbook(caminho, data_only=False)
        valores = load_workbook(caminho, data_only=True)
        nomes = self._wb.sheetnames
        if aba and aba not in nomes:
            raise ValueError(f"aba '{aba}' nao existe. Abas disponiveis: {nomes}")
        self.aba = aba or self._escolher_aba(valores)
        self._ws = self._wb[self.aba]
        ws_valores = valores[self.aba]

        self.linhas = []
        for linha_calc, linha_bruta in zip(
            ws_valores.iter_rows(values_only=True), self._ws.iter_rows(values_only=True)
        ):
            # Celula de formula sem valor em cache volta como None; nesse caso
            # fica a formula bruta, que o parser de dinheiro ignora com clareza.
            self.linhas.append(
                [calc if calc is not None else bruto for calc, bruto in zip(linha_calc, linha_bruta)]
            )
        self._largura = max((len(l) for l in self.linhas), default=0)
        self._validacoes: dict[tuple[int, int], list[str]] | None = None

    @staticmethod
    def _escolher_aba(wb) -> str:
        """A aba util e a que tem mais celulas preenchidas de verdade.

        Contar celulas em vez de dimensoes importa: a aba "Ajuda" do export do
        Mercado Livre declara mil linhas para pintar o fundo, mas guarda so
        alguns paragrafos de instrucao. Abas ocultas ficam de fora.
        """
        visiveis = [n for n in wb.sheetnames if wb[n].sheet_state == "visible"] or wb.sheetnames
        melhor, melhor_nota = visiveis[0], -1
        for nome in visiveis:
            preenchidas = sum(
                1
                for linha in wb[nome].iter_rows(max_row=5000, values_only=True)
                for celula in linha
                if celula is not None and str(celula).strip() != ""
            )
            if preenchidas > melhor_nota:
                melhor, melhor_nota = nome, preenchidas
        return melhor

    def escrever(self, linha: int, coluna: int, valor: object) -> None:
        if isinstance(valor, Decimal):
            valor = float(valor)
        self._ws.cell(row=linha + 1, column=coluna + 1, value=valor)
        while len(self.linhas) <= linha:
            self.linhas.append([])
        celulas = self.linhas[linha]
        celulas.extend([None] * (coluna + 1 - len(celulas)))
        celulas[coluna] = valor

    def acrescentar_coluna(self, titulo: str, linha_cabecalho: int) -> int:
        idx = self._largura
        self._largura += 1
        self.escrever(linha_cabecalho, idx, titulo)
        # A coluna nova herda o visual do cabecalho existente.
        celula = self._ws.cell(row=linha_cabecalho + 1, column=idx + 1)
        modelo = self._ws.cell(row=linha_cabecalho + 1, column=1)
        celula.font = copy(modelo.font)
        celula.fill = copy(modelo.fill)
        celula.alignment = copy(modelo.alignment)
        return idx

    def cor_de_fundo(self, linha: int, coluna: int) -> str | None:
        celula = self._ws.cell(row=linha + 1, column=coluna + 1)
        preenchimento = celula.fill
        if not preenchimento or not preenchimento.patternType:
            return None
        cor = preenchimento.fgColor
        return str(cor.rgb) if cor and cor.type == "rgb" else None

    def opcoes_validacao(self, linha: int, coluna: int) -> list[str] | None:
        """Le a lista suspensa da celula (``"Aplicar proposta,Nao aplicar"``).

        O Mercado Livre usa vocabularios diferentes na mesma coluna conforme a
        linha, entao gravar a palavra errada invalida o reenvio.
        """
        if self._validacoes is None:
            self._validacoes = self._indexar_validacoes()
        return self._validacoes.get((linha + 1, coluna + 1))

    def _indexar_validacoes(self) -> dict[tuple[int, int], list[str]]:
        from openpyxl.utils import range_boundaries

        indice: dict[tuple[int, int], list[str]] = {}
        for dv in self._ws.data_validations.dataValidation:
            if dv.type != "list" or not dv.formula1:
                continue
            opcoes = [o.strip() for o in dv.formula1.strip('"').split(",") if o.strip()]
            if not opcoes:
                continue
            for intervalo in dv.sqref.ranges:
                col_i, lin_i, col_f, lin_f = range_boundaries(str(intervalo))
                for lin in range(lin_i, lin_f + 1):
                    for col in range(col_i, col_f + 1):
                        indice[(lin, col)] = opcoes
        return indice

    def salvar(self, destino: str | Path) -> Path:
        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        self._wb.save(destino)
        return destino


class TabelaTexto(Tabela):
    def __init__(self, caminho: Path):
        self.caminho = caminho
        texto = _decodificar(caminho.read_bytes())
        self.dialeto = _dialeto(texto[:8192])
        self.linhas = [list(l) for l in csv.reader(io.StringIO(texto), dialect=self.dialeto)]
        self._largura = max((len(l) for l in self.linhas), default=0)

    def escrever(self, linha: int, coluna: int, valor: object) -> None:
        while len(self.linhas) <= linha:
            self.linhas.append([])
        celulas = self.linhas[linha]
        celulas.extend([""] * (coluna + 1 - len(celulas)))
        celulas[coluna] = "" if valor is None else str(valor)
        self._largura = max(self._largura, coluna + 1)

    def acrescentar_coluna(self, titulo: str, linha_cabecalho: int) -> int:
        idx = self._largura
        self.escrever(linha_cabecalho, idx, titulo)
        return idx

    def salvar(self, destino: str | Path) -> Path:
        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        with destino.open("w", encoding="utf-8-sig", newline="") as saida:
            escritor = csv.writer(saida, dialect=self.dialeto)
            for linha in self.linhas:
                escritor.writerow([("" if c is None else c) for c in linha])
        return destino


def abrir(caminho: str | Path, aba: str | None = None) -> Tabela:
    """Abre a planilha para leitura e escrita do resultado."""
    caminho = Path(caminho)
    if not caminho.exists():
        raise FileNotFoundError(f"planilha nao encontrada: {caminho}")
    sufixo = _validar_extensao(caminho)
    return TabelaExcel(caminho, aba) if sufixo in EXTENSOES_EXCEL else TabelaTexto(caminho)


def ler_linhas(caminho: str | Path, aba: str | None = None) -> list[list[object]]:
    """Le apenas os valores - usado pela tabela de custos."""
    return abrir(caminho, aba).linhas


def abas(caminho: str | Path) -> list[str]:
    """Nomes das abas de um Excel (lista vazia para arquivos de texto)."""
    caminho = Path(caminho)
    if caminho.suffix.lower() not in EXTENSOES_EXCEL:
        return []
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()
