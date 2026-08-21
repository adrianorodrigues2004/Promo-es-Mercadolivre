"""Gera uma planilha de promocoes no formato do export do Mercado Livre.

Serve para testar o motor ponta a ponta sem depender de um export real: usa os
MLBs e precos da tabela de custos e simula os descontos sugeridos de campanha.
"""

from __future__ import annotations

import argparse
import random
import sys
import warnings
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from promoml.custos import carregar_custos

CABECALHO = [
    "Código do anúncio",
    "Título do anúncio",
    "Campanha",
    "Preço sem desconto",
    "Preço com desconto sugerido",
    "Desconto sugerido (%)",
    "Preço mínimo",
    "Participar da promoção?",
]


def gerar(origem: Path, destino: Path, limite: int, semente: int) -> Path:
    warnings.filterwarnings("ignore")
    sorteio = random.Random(semente)
    tabela = carregar_custos(origem)

    wb = Workbook()
    ws = wb.active
    ws.title = "Promoções"
    ws.append(["Planilha de promoções - Mercado Livre"])
    ws.append(["Preencha o preço com desconto e marque SIM para participar."])
    ws.append([])
    ws.append(CABECALHO)
    for celula in ws[4]:
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="FFF2CC")

    escritas = 0
    for item in tabela.itens:
        if not item.id or item.preco_atual is None or item.preco_atual <= 0:
            continue
        preco = Decimal(item.preco_atual).quantize(Decimal("0.01"))
        desconto = sorteio.choice([5, 8, 10, 12, 15, 18, 20, 25, 30, 35, 40])
        sugerido = (preco * (Decimal(100 - desconto) / 100)).quantize(Decimal("0.01"))
        minimo = (preco * Decimal("0.45")).quantize(Decimal("0.01"))
        ws.append(
            [
                f"MLB{item.id}",
                item.titulo,
                sorteio.choice(["Ofertas do Dia", "Semana do Consumidor", "Queima de Estoque"]),
                float(preco),
                float(sugerido),
                desconto,
                float(minimo),
                None,
            ]
        )
        escritas += 1
        if escritas >= limite:
            break

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--custos", default="exemplos/custos_usuario.xlsx", type=Path)
    p.add_argument("--destino", default="exemplos/promocoes_ml.xlsx", type=Path)
    p.add_argument("--limite", type=int, default=300)
    p.add_argument("--semente", type=int, default=42)
    args = p.parse_args()
    caminho = gerar(args.custos, args.destino, args.limite, args.semente)
    print(f"planilha de exemplo gerada em {caminho}")


if __name__ == "__main__":
    main()
