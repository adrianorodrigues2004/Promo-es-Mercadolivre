"""Decisao ponta a ponta sobre um export do Mercado Livre."""

from decimal import Decimal

import pytest
from openpyxl import load_workbook

from promoml.custos import carregar_custos
from promoml.motor import ACEITO, CONTRAPROPOSTA, SEM_CUSTO, SEM_MARGEM_FIXO, processar
from promoml.regras import Regras, regras_de_dict


@pytest.fixture
def rodada(promo_ml_xlsx, custos_xlsx):
    return processar(promo_ml_xlsx, carregar_custos(custos_xlsx), Regras())


def por_mlb(rodada, mlb):
    return next(d for d in rodada.decisoes if d.id == mlb)


def test_reconhece_o_export_nativo_e_pula_as_instrucoes(rodada):
    assert rodada.perfil.nome == "mercado_livre"
    assert rodada.resumo.total == 5           # 5 anuncios, nenhuma linha de instrucao
    assert rodada.perfil.primeira_linha == 5  # dados comecam na linha 6


def test_escolhe_a_aba_de_dados_e_nao_a_de_ajuda(rodada):
    assert rodada.tabela.aba == "Promoções"


def test_aceita_a_proposta_quando_sobra_margem(rodada):
    decisao = por_mlb(rodada, "MLB4000000001")
    assert decisao.participar
    assert decisao.motivo == ACEITO
    assert decisao.preco_aplicado == decisao.preco_proposto
    assert decisao.margem_pct >= Decimal("0.05")


def test_contrapropoe_quando_o_desconto_derruba_a_margem(rodada):
    decisao = por_mlb(rodada, "MLB4000000002")
    assert decisao.participar
    assert decisao.motivo == CONTRAPROPOSTA
    assert decisao.preco_aplicado > decisao.preco_proposto
    assert decisao.preco_aplicado >= decisao.piso
    assert decisao.margem_pct >= Decimal("0.05")


def test_toda_participacao_respeita_o_piso_de_margem(rodada):
    for decisao in rodada.decisoes:
        if decisao.participar:
            assert decisao.preco_aplicado >= decisao.piso, decisao.titulo
            assert decisao.margem_pct >= Decimal("0.05"), decisao.titulo


def test_sai_da_promocao_fechada_que_nao_da_margem(rodada):
    """Linha ja participando, preco travado: a unica saida e nao participar."""
    decisao = por_mlb(rodada, "MLB4000000003")
    assert not decisao.participar
    assert not decisao.preco_editavel
    assert decisao.motivo == SEM_MARGEM_FIXO


def test_anuncio_sem_custo_fica_de_fora_e_vira_pendencia(rodada):
    decisao = por_mlb(rodada, "MLB9999999999")
    assert not decisao.participar
    assert decisao.motivo == SEM_CUSTO
    assert decisao in rodada.pendencias


def test_custo_casado_pelo_preco_quando_o_mlb_e_desconhecido(rodada):
    decisao = por_mlb(rodada, "MLB8888888888")
    assert decisao.custo_via == "preco atual"
    assert decisao.custo == Decimal("100")


def test_grava_a_palavra_do_dropdown_de_cada_linha(rodada, tmp_path):
    """"Aplicar proposta" numa linha e "Participar" na outra - trocar invalida."""
    destino = rodada.salvar(tmp_path / "saida.xlsx")
    ws = load_workbook(destino)["Promoções"]
    acoes = {ws.cell(row=i, column=2).value: ws.cell(row=i, column=14).value for i in range(6, 11)}
    assert acoes["MLB4000000001"] == "Aplicar proposta"
    assert acoes["MLB4000000002"] == "Aplicar proposta"
    assert acoes["MLB4000000003"] == "Não participar"   # vocabulario da linha fechada
    assert acoes["MLB9999999999"] == "Não aplicar"


def test_preco_e_desconto_gravados_se_correspondem(rodada, tmp_path):
    """O ML exige que o preco final case com o percentual que o acompanha."""
    destino = rodada.salvar(tmp_path / "saida.xlsx")
    ws = load_workbook(destino)["Promoções"]
    reescritas = [d.linha for d in rodada.decisoes if d.houve_contraproposta]
    assert reescritas, "nenhuma contraproposta para conferir"
    for linha in reescritas:
        original = ws.cell(row=linha, column=6).value
        desconto = ws.cell(row=linha, column=10).value
        final = ws.cell(row=linha, column=11).value
        assert abs(final - round(original * (1 - desconto / 100), 2)) <= 0.01


def test_nao_mexe_no_preco_das_linhas_fechadas(promo_ml_xlsx, custos_xlsx, tmp_path):
    antes = load_workbook(promo_ml_xlsx)["Promoções"]
    preco_antes = antes.cell(row=8, column=11).value

    rodada = processar(promo_ml_xlsx, carregar_custos(custos_xlsx), Regras())
    depois = load_workbook(rodada.salvar(tmp_path / "saida.xlsx"))["Promoções"]
    assert depois.cell(row=8, column=11).value == preco_antes


def test_planilha_original_nao_e_alterada(promo_ml_xlsx, custos_xlsx, tmp_path):
    antes = promo_ml_xlsx.read_bytes()
    processar(promo_ml_xlsx, carregar_custos(custos_xlsx), Regras()).salvar(tmp_path / "s.xlsx")
    assert promo_ml_xlsx.read_bytes() == antes


def test_colunas_de_diagnostico_saem_no_fim(rodada, tmp_path):
    ws = load_workbook(rodada.salvar(tmp_path / "saida.xlsx"))["Promoções"]
    cabecalhos = [c.value for c in ws[1]]
    assert "PROMO Motivo" in cabecalhos
    assert "PROMO Margem %" in cabecalhos
    assert cabecalhos.index("PROMO Participar") > 25   # depois das colunas do ML


def test_estrategia_so_proposta_nunca_mexe_no_preco(promo_ml_xlsx, custos_xlsx):
    regras = regras_de_dict({"estrategia": "so_proposta"})
    rodada = processar(promo_ml_xlsx, carregar_custos(custos_xlsx), regras)
    assert not any(d.houve_contraproposta for d in rodada.decisoes)
    assert not por_mlb(rodada, "MLB4000000002").participar


def test_estrategia_maior_desconto_desce_ate_o_piso(promo_ml_xlsx, custos_xlsx):
    regras = regras_de_dict({"estrategia": "maior_desconto"})
    rodada = processar(promo_ml_xlsx, carregar_custos(custos_xlsx), regras)
    decisao = por_mlb(rodada, "MLB4000000001")
    assert decisao.preco_aplicado < decisao.preco_proposto   # desconto maior que o proposto
    assert decisao.preco_aplicado >= decisao.piso


def test_margem_mais_exigente_reduz_a_participacao(promo_ml_xlsx, custos_xlsx):
    custos = carregar_custos(custos_xlsx)
    frouxa = processar(promo_ml_xlsx, custos, regras_de_dict({"margem_min_pct": "5%"}))
    apertada = processar(promo_ml_xlsx, custos, regras_de_dict({"margem_min_pct": "60%"}))
    assert apertada.resumo.participando < frouxa.resumo.participando


def test_desconto_minimo_evita_contraproposta_irrisoria(promo_ml_xlsx, custos_xlsx):
    """Uma contraproposta abaixo do minimo da campanha so suja o reenvio."""
    custos = carregar_custos(custos_xlsx)
    padrao = processar(promo_ml_xlsx, custos, Regras())
    exigente = processar(promo_ml_xlsx, custos, regras_de_dict({"desconto_min_padrao_pct": "35%"}))

    antes = por_mlb(padrao, "MLB4000000002")
    assert antes.houve_contraproposta and antes.desconto_aplicado_pct < 35
    assert not por_mlb(exigente, "MLB4000000002").participar
    for decisao in exigente.decisoes:
        if decisao.houve_contraproposta:
            assert decisao.desconto_aplicado_pct >= 35


def test_resumo_soma_o_que_foi_decidido(rodada):
    s = rodada.resumo
    assert s.total == s.participando + s.recusados
    assert sum(s.motivos.values()) == s.total
    assert s.lucro_previsto == sum(
        d.lucro for d in rodada.decisoes if d.participar
    )
