"""Fixtures sinteticas que imitam a estrutura dos arquivos reais.

Os testes nao dependem de nenhuma planilha de verdade: o export do Mercado
Livre e a tabela de precificacao sao reconstruidos aqui com as mesmas
particularidades que importam - linhas de instrucao antes dos dados, dois
vocabularios de acao na mesma coluna e celulas coloridas marcando o que pode
ser alterado.
"""

from __future__ import annotations

import sys
import warnings
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
warnings.filterwarnings("ignore", module="openpyxl")

from promoml.perfis import COR_EDITAVEL  # noqa: E402

CINZA = "FFF3F3F3"

CABECALHO_TECNICO = [
    "TITLE", "ITEM_ID", "SKU", "PROMO_NAME", "DATE", "ORIGINAL_PRICE", "SELLER_AMOUNT",
    "SELLER_PERCENTAGE", "MELI_AMOUNT", "DISCOUNT_PERCENTAGE", "FINAL_PRICE", "NEW_RECEIVES",
    "STATUS", "ACTION", "ERRORS", "ENTITY_ID", "TARGET_TYPE", "SUGGESTED_ACTION", "PROMO_ID",
    "PROMO_TYPE", "PROMO_SUB_TYPE", "START_DATE", "FINISH_DATE", "TASK_RESOLUTION_RULES",
    "CANDIDATE_ID", "SALE_FEE",
]


@pytest.fixture
def custos_xlsx(tmp_path: Path) -> Path:
    """Tabela de precificacao no mesmo formato da planilha real."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Mercado Livre"
    ws.append(
        ["Produto", "Código", "Custo", "Frete ", "Imposto", "Taxa", "Margem Desejada",
         "Anunciar", "Preço Atual", "Promoção (%)", "Promoção", "Cupom (%)", "Ribait",
         "Preço Final", "Lucro Líquido ", "Margem Líquida"]
    )
    # produto, codigo, custo, frete, imposto, taxa, margem desejada, anunciar, preco atual
    ws.append(["MESA FOLGADA", "4000000001", 100, 20, 11.5, 11.5, 30, 0, 500])
    ws.append(["MESA APERTADA", "4000000002", 380, 50, 11.5, 16.5, 30, 0, 900])
    ws.append(["BANCO BARATO", "4000000003", 40, 0, 11.5, 11.5, 30, 0, 100])
    ws.append(["ITEM SO POR PRECO", "", 100, 10, 11.5, 11.5, 30, 0, 777.77])
    destino = tmp_path / "custos.xlsx"
    wb.save(destino)
    return destino


def _linha_promo(titulo, mlb, sku, original, final, desconto, meli, recebe, alvo, subtipo):
    linha = [None] * 26
    linha[0], linha[1], linha[2] = titulo, mlb, sku
    linha[3], linha[4] = "Campanha de teste", "De 1 a 30 de Janeiro"
    linha[5], linha[6] = original, round(original - final - meli, 2)
    linha[8], linha[9], linha[10] = meli, desconto, final
    linha[11] = f"$\xa0{recebe}\n(Inclui uma redução nas suas tarifas de venda de R${meli})"
    linha[12] = "Nova proposta" if alvo == "OPTINEABLE" else "Participando"
    linha[13] = "Aplicar proposta" if alvo == "OPTINEABLE" else "Participar"
    linha[16], linha[19], linha[20] = alvo, "SMART_CAMPAIGN", subtipo
    linha[25] = meli
    return linha


@pytest.fixture
def promo_ml_xlsx(tmp_path: Path) -> Path:
    """Export de campanha do ML com as cinco linhas de cabecalho e instrucao."""
    wb = Workbook()
    ajuda = wb.active
    ajuda.title = "Ajuda"
    ajuda["B2"] = "Gerencie suas promoções"
    ajuda.cell(row=900, column=20, value=None)  # muitas linhas, poucas preenchidas

    ws = wb.create_sheet("Promoções")
    ws.append(CABECALHO_TECNICO)
    ws.append(["Título do anúncio", "Número do anúncio", "SKU", "Promoção", "Vigência",
               "Preço original", "Desconto"] + [None] * 19)
    ws.append([None] * 6 + ["Valor por sua conta", "Porcentagem por sua conta",
                            "Redução nas suas tarifas de venda", "Desconto total"] + [None] * 16)
    ws.append(["Não altere esta coluna"] * 9 + ["Você só pode alterar a porcentagem"] * 2
              + ["Não altere esta coluna", "Não altere esta coluna", "", "Não altere esta coluna"]
              + [None] * 11)
    ws.append([None] * 26)

    linhas = [
        # sobra margem com o desconto proposto -> aceita como esta
        _linha_promo("Mesa Folgada Grande", "MLB4000000001", "SKU-A", 500, 400, 20, 5, 350,
                     "OPTINEABLE", "ON_DEMAND"),
        # desconto proposto derruba a margem, mas da para renegociar
        _linha_promo("Mesa Apertada Media", "MLB4000000002", "SKU-B", 900, 630, 30, 6, 482.05,
                     "OPTINEABLE", "COFINANCED"),
        # ja participando: preco fechado, so da para sair
        _linha_promo("Banco Barato Fixo", "MLB4000000003", "SKU-C", 100, 60, 40, 2, 50,
                     "OFFER", "ON_DEMAND"),
        # sem custo cadastrado
        _linha_promo("Produto Desconhecido", "MLB9999999999", "SKU-Z", 300, 250, 17, 3, 200,
                     "OPTINEABLE", "ON_DEMAND"),
        # casa pelo preco atual, porque o MLB nao esta na tabela de custos
        _linha_promo("Item So Por Preco", "MLB8888888888", "SKU-P", 777.77, 700, 10, 4, 600,
                     "OPTINEABLE", "ON_DEMAND"),
    ]
    for linha in linhas:
        ws.append(linha)

    azul, cinza = PatternFill("solid", fgColor=COR_EDITAVEL), PatternFill("solid", fgColor=CINZA)
    editaveis, fechadas = [], []
    for i, linha in enumerate(linhas, start=6):
        alvo = linha[16]
        preenchimento = azul if alvo == "OPTINEABLE" else cinza
        for coluna in (10, 11):
            ws.cell(row=i, column=coluna).fill = preenchimento
        (editaveis if alvo == "OPTINEABLE" else fechadas).append(f"N{i}")

    if editaveis:
        dv = DataValidation(type="list", formula1='"Aplicar proposta,Não aplicar"')
        ws.add_data_validation(dv)
        for ref in editaveis:
            dv.add(ref)
    if fechadas:
        dv = DataValidation(type="list", formula1='"Participar,Não participar"')
        ws.add_data_validation(dv)
        for ref in fechadas:
            dv.add(ref)

    destino = tmp_path / "promocoes.xlsx"
    wb.save(destino)
    return destino
