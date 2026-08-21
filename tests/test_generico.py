"""Planilhas que nao sao o export nativo do ML caem no perfil generico."""

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from promoml.custos import carregar_custos
from promoml.motor import processar
from promoml.regras import Regras, regras_de_dict


@pytest.fixture
def planilha_simples(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["Planilha de promoções"])          # titulo solto
    ws.append(["Não altere as colunas bloqueadas"])
    ws.append([])
    ws.append(["Código do anúncio", "Título do anúncio", "Preço sem desconto",
               "Preço com desconto sugerido", "Preço mínimo", "Participar da promoção?"])
    ws.append(["MLB4000000001", "Mesa Folgada", 500, 400, 200, None])
    ws.append(["MLB4000000003", "Banco Barato", 100, 55, 40, None])
    destino = tmp_path / "simples.xlsx"
    wb.save(destino)
    return destino


@pytest.fixture
def csv_simples(tmp_path: Path) -> Path:
    destino = tmp_path / "simples.csv"
    destino.write_text(
        "Código do anúncio;Título;Preço sem desconto;Preço promocional;Participar\n"
        "MLB4000000001;Mesa Folgada;500,00;400,00;\n"
        "MLB4000000003;Banco Barato;100,00;55,00;\n",
        encoding="utf-8",
    )
    return destino


def test_perfil_generico_encontra_cabecalho_e_decide(planilha_simples, custos_xlsx):
    rodada = processar(planilha_simples, carregar_custos(custos_xlsx), Regras())
    assert rodada.perfil.nome == "generico"
    assert rodada.resumo.total == 2
    assert {d.id for d in rodada.decisoes} == {"MLB4000000001", "MLB4000000003"}


def test_generico_grava_sim_nao_e_preco(planilha_simples, custos_xlsx, tmp_path):
    rodada = processar(planilha_simples, carregar_custos(custos_xlsx), Regras())
    ws = load_workbook(rodada.salvar(tmp_path / "saida.xlsx")).active
    decisoes = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=6).value for i in (5, 6)}
    assert set(decisoes.values()) <= {"SIM", "NAO"}
    for decisao in rodada.decisoes:
        if decisao.participar:
            assert decisao.margem_pct >= Decimal("0.05")


def test_generico_contrapoe_dentro_do_preco_atual(planilha_simples, custos_xlsx):
    """Sem coluna de desconto, a contraproposta nunca passa do preco cheio."""
    rodada = processar(planilha_simples, carregar_custos(custos_xlsx), Regras())
    for decisao in rodada.decisoes:
        if decisao.participar:
            assert decisao.preco_aplicado <= decisao.preco_atual


def test_csv_e_lido_e_reescrito(csv_simples, custos_xlsx, tmp_path):
    rodada = processar(csv_simples, carregar_custos(custos_xlsx), Regras())
    assert rodada.resumo.total == 2
    destino = rodada.salvar(tmp_path / "saida.csv")
    texto = destino.read_text(encoding="utf-8-sig")
    assert "PROMO Motivo" in texto
    assert texto.count("\n") >= 3


def test_arredondamento_configuravel(planilha_simples, custos_xlsx):
    custos = carregar_custos(custos_xlsx)
    regras = regras_de_dict({"arredondamento": "terminacao_90", "estrategia": "maior_desconto"})
    rodada = processar(planilha_simples, custos, regras)
    for decisao in rodada.decisoes:
        if decisao.participar:
            assert str(decisao.preco_aplicado).endswith(".90")


def test_planilha_sem_coluna_de_preco_falha_com_mensagem_util(tmp_path, custos_xlsx):
    wb = Workbook()
    wb.active.append(["Código do anúncio", "Título do anúncio"])
    wb.active.append(["MLB1", "Mesa"])
    caminho = tmp_path / "sem_preco.xlsx"
    wb.save(caminho)
    with pytest.raises(ValueError, match="coluna de preco"):
        processar(caminho, carregar_custos(custos_xlsx), Regras())


def test_xls_antigo_e_recusado_com_orientacao(tmp_path, custos_xlsx):
    caminho = tmp_path / "antigo.xls"
    caminho.write_bytes(b"nao importa")
    with pytest.raises(ValueError, match="salve como .xlsx"):
        processar(caminho, carregar_custos(custos_xlsx), Regras())
