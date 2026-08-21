"""Tabela de custos: leitura das colunas e casamento com o anuncio."""

from decimal import Decimal

import pytest

from promoml.custos import carregar_custos
from promoml.regras import Regras


def test_le_as_colunas_da_planilha_de_precificacao(custos_xlsx):
    tabela = carregar_custos(custos_xlsx)
    assert tabela.campos["custo"] == "Custo"
    assert tabela.campos["comissao_pct"] == "Taxa"
    assert tabela.campos["frete"] == "Frete "        # com espaco sobrando
    assert tabela.campos["id"] == "Código"
    assert tabela.campos["preco_atual"] == "Preço Atual"


def test_margem_desejada_nao_vira_margem_minima(custos_xlsx):
    """"Margem Desejada" e meta comercial, nao regra de corte da promocao."""
    tabela = carregar_custos(custos_xlsx)
    assert "margem_min_pct" not in tabela.campos
    custo, _ = tabela.buscar(id_anuncio="MLB4000000001")
    assert "margem_min_pct" not in custo.sobrescritas


def test_casa_mlb_com_e_sem_prefixo(custos_xlsx):
    """A planilha de custos guarda 4000000001; o ML manda MLB4000000001."""
    tabela = carregar_custos(custos_xlsx)
    custo, via = tabela.buscar(id_anuncio="MLB4000000001")
    assert custo.titulo == "MESA FOLGADA"
    assert via == "MLB (numero)"


def test_casa_pelo_preco_atual_quando_nao_ha_mlb(custos_xlsx):
    tabela = carregar_custos(custos_xlsx)
    custo, via = tabela.buscar(id_anuncio="MLB0000000000", preco_atual=Decimal("777.77"))
    assert custo.titulo == "ITEM SO POR PRECO"
    assert via == "preco atual"


def test_preco_ambiguo_nao_chuta_um_custo(custos_xlsx, tmp_path):
    """Dois produtos com o mesmo preco e custos diferentes: melhor nao casar."""
    from openpyxl import load_workbook

    wb = load_workbook(custos_xlsx)
    wb["Mercado Livre"].append(["OUTRO ITEM", "", 999, 10, 11.5, 11.5, 30, 0, 777.77])
    destino = tmp_path / "ambiguo.xlsx"
    wb.save(destino)

    tabela = carregar_custos(destino)
    custo, via = tabela.buscar(preco_atual=Decimal("777.77"), titulo="Nome que nao ajuda")
    assert custo is None and via == ""


def test_titulo_desempata_precos_iguais(custos_xlsx, tmp_path):
    from openpyxl import load_workbook

    wb = load_workbook(custos_xlsx)
    wb["Mercado Livre"].append(["BANCO CARO", "", 999, 10, 11.5, 11.5, 30, 0, 777.77])
    destino = tmp_path / "desempate.xlsx"
    wb.save(destino)

    tabela = carregar_custos(destino)
    custo, via = tabela.buscar(preco_atual=Decimal("777.77"), titulo="Banco Caro De Madeira")
    assert custo.titulo == "BANCO CARO"
    assert via == "preco + titulo"


def test_metodos_desabilitados_sao_respeitados(custos_xlsx):
    tabela = carregar_custos(custos_xlsx)
    custo, _ = tabela.buscar(preco_atual=Decimal("777.77"), metodos=("mlb", "sku"))
    assert custo is None


def test_encargos_por_anuncio_sobrescrevem_as_regras_globais(custos_xlsx):
    """Classico (11,5%) e Premium (16,5%) convivem na mesma rodada."""
    tabela = carregar_custos(custos_xlsx)
    regras = Regras()
    classico, _ = tabela.buscar(id_anuncio="4000000001")
    premium, _ = tabela.buscar(id_anuncio="4000000002")
    assert classico.encargos(regras).comissao_pct == Decimal("0.115")
    assert premium.encargos(regras).comissao_pct == Decimal("0.165")
    assert premium.encargos(regras).frete == Decimal("50")


def test_arquivo_sem_coluna_de_custo_falha_com_mensagem_util(tmp_path):
    caminho = tmp_path / "ruim.csv"
    caminho.write_text("produto;preco\nmesa;10\n", encoding="utf-8")
    with pytest.raises(ValueError, match="coluna de custo"):
        carregar_custos(caminho)


def test_varias_tabelas_se_completam(custos_xlsx, tmp_path):
    """O arquivo de pendencias preenchido soma a tabela principal, nao a troca."""
    complemento = tmp_path / "pendencias.csv"
    complemento.write_text(
        "mlb;sku;titulo;preco_atual;preco_proposto_ml;custo\n"
        "MLB9999999999;SKU-Z;Produto Desconhecido;300,00;250,00;90,00\n",
        encoding="utf-8",
    )
    tabela = carregar_custos([custos_xlsx, complemento])

    novo, _ = tabela.buscar(id_anuncio="MLB9999999999")
    antigo, _ = tabela.buscar(id_anuncio="MLB4000000001")
    assert novo.custo == Decimal("90.00")     # veio do complemento
    assert antigo.custo == Decimal("100")     # continua valendo


def test_tabela_mais_nova_corrige_a_anterior(custos_xlsx, tmp_path):
    correcao = tmp_path / "correcao.csv"
    correcao.write_text("mlb;custo\n4000000001;123,45\n", encoding="utf-8")
    tabela = carregar_custos([custos_xlsx, correcao])
    custo, _ = tabela.buscar(id_anuncio="MLB4000000001")
    assert custo.custo == Decimal("123.45")


def test_lista_vazia_de_tabelas_e_recusada():
    with pytest.raises(ValueError, match="nenhuma tabela"):
        carregar_custos([])


def test_pasta_de_extras_entra_sozinha(custos_xlsx, tmp_path):
    """Arquivo guardado na pasta de extras nao precisa ser reenviado."""
    from promoml.custos import descobrir_tabelas

    extras = tmp_path / "custos-extras"
    extras.mkdir()
    (extras / "2026-08-pendencias.csv").write_text("mlb;custo\n5000000001;10,00\n", encoding="utf-8")

    caminhos = descobrir_tabelas(custos_xlsx, extras)
    assert [c.name for c in caminhos] == ["custos.xlsx", "2026-08-pendencias.csv"]

    tabela = carregar_custos(caminhos)
    novo, _ = tabela.buscar(id_anuncio="MLB5000000001")
    assert novo.custo == Decimal("10.00")


def test_extras_mais_recente_por_nome_corrige_o_anterior(custos_xlsx, tmp_path):
    from promoml.custos import descobrir_tabelas

    extras = tmp_path / "custos-extras"
    extras.mkdir()
    (extras / "2026-08-pendencias.csv").write_text("mlb;custo\n5000000001;10,00\n", encoding="utf-8")
    (extras / "2026-09-pendencias.csv").write_text("mlb;custo\n5000000001;99,00\n", encoding="utf-8")

    tabela = carregar_custos(descobrir_tabelas(custos_xlsx, extras))
    custo, _ = tabela.buscar(id_anuncio="MLB5000000001")
    assert custo.custo == Decimal("99.00")


def test_pasta_de_extras_ignora_instrucoes_e_temporarios(custos_xlsx, tmp_path):
    from promoml.custos import descobrir_tabelas

    extras = tmp_path / "custos-extras"
    extras.mkdir()
    (extras / "LEIA-ME.txt").write_text("Guarde aqui os custos que faltavam.", encoding="utf-8")
    (extras / "~$rascunho.xlsx").write_text("lixo do Excel", encoding="utf-8")
    (extras / ".DS_Store").write_text("lixo do Mac", encoding="utf-8")
    (extras / "pendencias.csv").write_text("mlb;custo\n5000000001;10,00\n", encoding="utf-8")

    assert [c.name for c in descobrir_tabelas(custos_xlsx, extras)[1:]] == ["pendencias.csv"]


def test_complemento_ilegivel_vira_aviso_e_nao_derruba_a_rodada(custos_xlsx, tmp_path):
    ruim = tmp_path / "sem_custo.csv"
    ruim.write_text("produto;preco\nmesa;10\n", encoding="utf-8")

    tabela = carregar_custos([custos_xlsx, ruim])
    assert tabela.buscar(id_anuncio="MLB4000000001")[0].custo == Decimal("100")
    assert any("sem_custo.csv ignorado" in aviso for aviso in tabela.avisos)


def test_tabela_principal_ilegivel_continua_sendo_erro(tmp_path, custos_xlsx):
    ruim = tmp_path / "sem_custo.csv"
    ruim.write_text("produto;preco\nmesa;10\n", encoding="utf-8")
    with pytest.raises(ValueError, match="coluna de custo"):
        carregar_custos([ruim, custos_xlsx])


def test_pasta_de_extras_inexistente_nao_atrapalha(custos_xlsx, tmp_path):
    from promoml.custos import descobrir_tabelas

    assert descobrir_tabelas(custos_xlsx, tmp_path / "nao-existe") == [custos_xlsx]
